# 文件：backend/app/providers/docs.py
# 作用：PDF/Word 解析（P19）：按 config.DOCS_BACKEND 排解析后端顺序，逐个兜底，全失败给可读报错
# 阶段：F9 作业与文档（兼 P19）
# 依赖：markitdown、docling（可选）、pypdf、python-docx、app/core/config.py
from __future__ import annotations

import importlib.util
from pathlib import Path

from app.core import config

SUFFIXES = (".pdf", ".docx")


class DocsError(RuntimeError):
    """文档解析失败：文件坏了、有口令保护、或一个解析后端都没装。"""


def _available(name: str) -> bool:
    """依赖装了没：解析后端是可选能力，装不上就退到下一个，不硬崩。"""
    return importlib.util.find_spec(name) is not None


ENABLED = _available("markitdown") or _available("pypdf")


def extract(path: Path) -> tuple[str, str]:
    """解析一份 PDF/Word，返回 (正文, source_type)；全部后端都失败时抛 DocsError。"""
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise DocsError(f"只支持 {'/'.join(SUFFIXES)}，收到：{path.name}")
    errors: list[str] = []
    for name in _order():
        try:
            text = _BACKENDS[name](path)
        except DocsError as exc:
            errors.append(str(exc))
            continue
        except Exception as exc:  # noqa: BLE001  解析器五花八门的异常：换下一个后端，别把 500 抛给用户
            errors.append(f"{name}: {exc}")
            continue
        if text.strip():
            return text, suffix.lstrip(".")
        errors.append(f"{name}: 没解出正文")
    raise DocsError(f"解不出 {path.name} 的正文（{'；'.join(errors) or '没有可用的解析后端'}）")


def _order() -> list[str]:
    """解析顺序：默认轻后端在前（无模型下载）；DOCS_BACKEND=docling 时把重后端提到最前。"""
    heavy_first = config.DOCS_BACKEND == "docling"
    order = ["docling", "markitdown"] if heavy_first else ["markitdown", "docling"]
    order = [name for name in order if name in _installed()]
    order.append("python")  # 兜底后端总在最后（pypdf / python-docx，本仓库的必装件）
    return order


def _installed() -> list[str]:
    """装了哪些可插拔后端。"""
    return [name for name in ("markitdown", "docling") if _available(name)]


def _by_markitdown(path: Path) -> str:
    """MarkItDown（轻）：docx/pdf/pptx/xlsx 一个入口。"""
    from markitdown import MarkItDown

    return MarkItDown().convert(str(path)).text_content


def _by_docling(path: Path) -> str:
    """Docling（重）：版面理解较好，首次运行会下模型，所以默认不排在最前。"""
    from docling.document_converter import DocumentConverter

    return DocumentConverter().convert(str(path)).document.export_to_markdown()


def _by_python(path: Path) -> str:
    """兜底：PDF 走 pypdf，Word 走 python-docx（含表格文字）。"""
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            raise DocsError(f"{path.name} 有口令保护（加密 PDF），先解密再导入")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


_BACKENDS = {"markitdown": _by_markitdown, "docling": _by_docling, "python": _by_python}
