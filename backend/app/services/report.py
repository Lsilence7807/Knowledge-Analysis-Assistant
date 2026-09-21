# 文件：backend/app/services/report.py
# 作用：报告导出（§4.2）：按 task_id 取问题/SQL/结论，重跑 SQL 拿结果表，导出 docx/pptx/md
# 阶段：F9 作业与文档（兼 P19）
# 依赖：docxtpl、python-pptx、jinja2、app/core/{config,db}.py、app/services/store.py
from __future__ import annotations

from datetime import UTC, datetime

from docxtpl import DocxTemplate
from jinja2 import Template
from pptx import Presentation
from pptx.util import Inches, Pt

from app.core import config, db
from app.services import store

FORMATS = ("docx", "pptx", "md")
EXPORT_ROWS = 200  # 报告里最多放多少行明细（与 config/tools.json 的单步上限同口径）
PPT_MAX_ROWS = 12  # 一页幻灯片放不下太多行，超出只放前 12 行

MD_TEMPLATE = """# {{ question }}

- 数据集：{{ dataset_id }}
- 生成时间：{{ generated_at }}

## SQL

    {{ sql }}

## 结论

{{ insight }}

## 结果表（前 {{ rows|length }} 行）

{% for line in table_lines %}{{ line }}
{% endfor %}"""


class ReportError(RuntimeError):
    """导出失败（格式不支持、任务没有 SQL、结果表重跑失败）。"""


class ReportNotFound(ReportError):
    """任务不存在，路由据此回 404。"""


def export(task_id: str, fmt: str) -> str:
    """导出报告，返回相对仓库根的 posix 路径（§4.2）。

    先写 `<名字>.part` 再原子改名：中途失败不留半个文件（P19 反例口径）。
    """
    fmt = str(fmt or "").strip().lower().lstrip(".")
    if fmt not in FORMATS:
        raise ReportError(f"只支持 {'/'.join(FORMATS)}，收到：{fmt or '空'}")
    task = store.get_task(task_id)
    if task is None:
        raise ReportNotFound(f"任务不存在：{task_id or '空'}")
    if not str(task.get("sql") or "").strip():
        raise ReportError("这次任务没有可导出的 SQL（当时降级了），换个成功的任务再导")
    doc = _document(task, _result(task))
    config.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    target = config.EXPORT_DIR / f"{task_id}.{fmt}"
    staging = target.with_name(target.name + ".part")
    try:
        _WRITERS[fmt](doc, staging)
        staging.replace(target)
    except Exception as exc:  # noqa: BLE001  写入失败统一成可读错误；临时文件在 finally 里清掉
        raise ReportError(f"导出失败：{exc}") from exc
    finally:
        if staging.exists():
            staging.unlink()
    try:
        return target.relative_to(config.BASE_DIR).as_posix()
    except ValueError:  # 导出目录被换到仓库外（测试或自定义 DATA_DIR）时给绝对 posix 路径
        return target.as_posix()


def _result(task: dict) -> dict:
    """重跑这次任务的 SQL 拿结果表：tasks 只存问题与 SQL，明细没落库（§4.4）。"""
    try:
        return db.exec_sql(str(task["sql"]), limit=EXPORT_ROWS)
    except db.SQLRejected as exc:
        raise ReportError(f"结果表重跑失败：{exc}") from exc


def _document(task: dict, table: dict) -> dict:
    """报告要渲染的字段：问题、SQL、结论与结果表（每行先拍成一行文本，模板不必关心列数）。"""
    header = [str(name) for name in table.get("columns") or []]
    rows = [" | ".join(str(cell) for cell in row) for row in table.get("rows") or []]
    insight = store.latest_insight(str(task.get("session_id") or ""), str(task.get("question") or ""))
    return {
        "question": task.get("question") or "",
        "dataset_id": task.get("dataset_id") or "",
        "sql": task.get("sql") or "",
        "insight": insight or "（这次没有留下结论摘要）",
        "header": header,
        "rows": rows,
        "table_lines": ([" | ".join(header)] if header else []) + rows,
        "generated_at": datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC"),
    }


def _write_md(doc: dict, path) -> None:
    """Markdown：Jinja2 直接渲染，没有模板文件。"""
    path.write_text(Template(MD_TEMPLATE).render(**doc), encoding="utf-8")


def _write_docx(doc: dict, path) -> None:
    """Word：docxtpl 渲染仓库里的模板（表格用 1 列表格逐行铺，避免列数与模板耦合）。"""
    template = DocxTemplate(str(config.REPORT_TEMPLATE))
    template.render(doc)
    template.save(str(path))


def _write_pptx(doc: dict, path) -> None:
    """PPT：python-pptx 现场搭一页（标题 + 结论 + 结果表），不用模板。"""
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = (doc["question"] or "数据分析报告")[:80]
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(9), Inches(1.4)).text_frame
    body.text = doc["insight"]
    sql_line = body.add_paragraph()
    sql_line.text = f"SQL：{doc['sql'][:200]}"
    for paragraph in body.paragraphs:
        paragraph.font.size = Pt(12)
    columns = max(len(doc["header"]), 1)
    lines = doc["rows"][:PPT_MAX_ROWS]
    table = slide.shapes.add_table(len(lines) + 1, columns, Inches(0.6), Inches(3.0), Inches(9), Inches(3.5)).table
    for index in range(columns):
        table.cell(0, index).text = doc["header"][index] if index < len(doc["header"]) else ""
    for row_index, line in enumerate(lines, start=1):
        for col_index in range(columns):
            cells = line.split(" | ")
            table.cell(row_index, col_index).text = cells[col_index] if col_index < len(cells) else ""
    deck.save(str(path))


_WRITERS = {"md": _write_md, "docx": _write_docx, "pptx": _write_pptx}
