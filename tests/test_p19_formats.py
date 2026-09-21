# 文件：tests/test_p19_formats.py
# 作用：F9/P19 验收：docx/pdf 导入后可检索（解析后端 + 兜底）、加密 PDF 与坏文件被拒且提示可读、
#       docx/pptx/md 导出能被重新打开且含表格与结论、导出失败不留半个文件
# 阶段：F9 作业与文档（兼 P19）
# 依赖：pytest、fastapi.testclient、python-docx、python-pptx、pypdf、app/core/config.py、app/main.py、
#       app/providers/{docs,kb}.py、app/services/{ingest,memory,report,store}.py
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pptx import Presentation
from pypdf import PdfReader, PdfWriter

from app.core import config
from app.main import create_app
from app.providers import docs, kb
from app.services import ingest, memory, report, store

QUESTION = "各地区的销售额合计是多少"
INSIGHT = "华东区销售额 1200 万元，环比下降 20%"
SAMPLE = "examples/demo_sales.csv"


@pytest.fixture()
def env(tmp_path, monkeypatch):
    """数据目录、知识库根目录与导出目录全指向临时目录；知识库能力显式打开（默认关）。"""
    monkeypatch.setattr(config, "DATA_DIR", tmp_path)
    monkeypatch.setattr(config, "SQLITE_PATH", tmp_path / "meta.sqlite")
    monkeypatch.setattr(config, "DUCKDB_PATH", tmp_path / "analytics.duckdb")
    monkeypatch.setattr(config, "LOCAL_SETTINGS", tmp_path / "local.json")
    monkeypatch.setattr(config, "EXPORT_DIR", tmp_path / "exports")
    monkeypatch.setattr(config, "ENABLE_AGENT", False)
    monkeypatch.setattr(kb, "ROOT_DIR", tmp_path)
    monkeypatch.setattr(kb, "ENABLED", True)
    (tmp_path / "kb").mkdir(parents=True, exist_ok=True)
    return tmp_path


@pytest.fixture()
def client(env):
    with TestClient(create_app()) as test_client:
        yield test_client


# ---- 夹具：一份 docx、一份最小 PDF、一份加密 PDF ----


def _write_docx(root: Path, name: str, text: str) -> Path:
    path = root / "kb" / name
    document = Document()
    document.add_paragraph(text)
    document.save(str(path))
    return path


def _write_pdf(root: Path, name: str, text: str) -> Path:
    """手搓一份最小 PDF：不引新依赖，pypdf / pdfminer 都能抽出这行字。"""
    path = root / "kb" / name
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def _write_encrypted_pdf(root: Path, name: str) -> Path:
    """加密 PDF：先造一份正常的，再用 pypdf 加口令。"""
    source = _write_pdf(root, name, "quarterly revenue drop 20 percent")
    reader = PdfReader(str(source))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("secret")
    with source.open("wb") as handle:
        writer.write(handle)
    return source


def _seed_task(dataset_id: str, sql: str) -> str:
    """造一条已完成的提问记录 + 一轮会话结论（导出报告的数据来源）。"""
    store.insert_task(
        {
            "id": "t_export",
            "dataset_id": dataset_id,
            "session_id": "s_export",
            "kind": "ask",
            "question": QUESTION,
            "sql": sql,
            "degraded": [],
        }
    )
    memory.add_turn("s_export", QUESTION, sql, ["地区", "销售额合计"], {"summary": INSIGHT}, False)
    return "t_export"


@pytest.fixture()
def seeded(env):
    """一份真实数据集 + 一条可导出的任务；SQL 直接算真的（导出会重跑它）。"""
    dataset = ingest.ingest_file(config.BASE_DIR / SAMPLE, "demo_sales.csv")
    sql = f"SELECT 地区, sum(销售额) AS 销售额合计 FROM {dataset['table']} GROUP BY 地区 ORDER BY 2 DESC"
    return {"dataset_id": dataset["dataset_id"], "task_id": _seed_task(dataset["dataset_id"], sql)}


# ---- 导入：PDF/Word 也能进知识库 ----


def test_docx_is_imported_and_searchable(client, env):
    """自动：docx 走解析后端入库，检索命中且带出处。"""
    _write_docx(env, "季度口径.docx", "华东区第三季度销售额环比下降百分之二十，口径是万元。" * 3)
    body = client.post("/kb/import", json={"path": "kb/季度口径.docx"}).json()
    assert body["skipped"] == [] and body["imported"][0]["chunks"] >= 1
    with store.connect() as conn:
        kind = conn.execute("SELECT source_type FROM kb_docs WHERE path = 'kb/季度口径.docx'").fetchone()[0]
    assert kind == "docx"
    hits = client.post("/kb/search", json={"query": "环比下降"}).json()["hits"]
    assert hits and hits[0]["source"].endswith("季度口径.docx")


def test_pdf_is_imported_and_searchable(client, env):
    """自动：PDF 走解析后端入库并可检索（pypdf 兜底路径也要通）。"""
    _write_pdf(env, "quarterly.pdf", "quarterly revenue drop 20 percent")
    body = client.post("/kb/import", json={"path": "kb/quarterly.pdf"}).json()
    assert body["skipped"] == []
    hits = client.post("/kb/search", json={"query": "revenue"}).json()["hits"]
    assert hits and hits[0]["source"].endswith("quarterly.pdf")


def test_encrypted_pdf_is_rejected_with_readable_reason(client, env):
    """反例：加密 PDF 被拒，理由是人话（不是 traceback）。"""
    _write_encrypted_pdf(env, "secret.pdf")
    response = client.post("/kb/import", json={"path": "kb/secret.pdf"})
    assert response.status_code == 400
    assert "口令" in response.json()["detail"] or "加密" in response.json()["detail"]


def test_directory_import_skips_broken_file(client, env):
    """反例：目录里坏掉的那份只跳过，别拖垮整批。"""
    _write_docx(env, "好的.docx", "华东区销售额环比下降百分之二十。" * 3)
    # 坏文件要做成「真 docx 被截断」：纯垃圾字节 markitdown 会当纯文本收下，不算坏
    broken = env / "kb" / "坏的.docx"
    hole = _write_docx(env, "坏的.docx", "坏文件" * 5)
    broken.write_bytes(hole.read_bytes()[:120])
    body = client.post("/kb/import", json={}).json()
    assert [item["path"] for item in body["imported"]] == ["kb/好的.docx"]
    assert [item["path"] for item in body["skipped"]] == ["kb/坏的.docx"]
    assert body["skipped"][0]["error"]


def test_large_docx_is_chunked(client, env):
    """自动：超过分块上限的 docx 切成多块（P19 手工项的前置）。"""
    _write_docx(env, "长文.docx", "华东区第三季度销售额环比下降百分之二十。" * 60)
    body = client.post("/kb/import", json={"path": "kb/长文.docx"}).json()
    assert body["imported"][0]["chunks"] > 1


def test_docs_extract_reports_source_type(env):
    """自动 + 反例：解析器直接调用时回 (正文, source_type)；后缀不支持时给 DocsError。"""
    pdf = _write_pdf(env, "plain.pdf", "plain pdf text")
    text, source_type = docs.extract(pdf)
    assert "plain pdf text" in text and source_type == "pdf"
    with pytest.raises(docs.DocsError):
        docs.extract(env / "kb" / "note.md")


# ---- 导出：三种格式都能被重新打开 ----


def test_export_md_has_conclusion_and_table(client, seeded):
    """自动：Markdown 里有结论与结果表行。"""
    body = client.post("/exports", json={"task_id": seeded["task_id"], "fmt": "md"}).json()
    text = (config.BASE_DIR / body["path"]).read_text(encoding="utf-8")
    assert INSIGHT in text and "地区 | 销售额合计" in text


def test_export_docx_reopens_with_table_and_conclusion(client, seeded):
    """自动：导出的 docx 能被 python-docx 重新打开，表格与结论都在。"""
    body = client.post("/exports", json={"task_id": seeded["task_id"], "fmt": "docx"}).json()
    document = Document(str(config.BASE_DIR / body["path"]))
    lines = [cell.text for row in document.tables[0].rows for cell in row.cells]
    assert lines[0].startswith("地区 | 销售额合计") and len(lines) > 1
    assert any(INSIGHT in paragraph.text for paragraph in document.paragraphs)


def test_export_pptx_reopens_with_table_and_conclusion(client, seeded):
    """自动：导出的 pptx 能被 python-pptx 重新打开，表格与结论都在。"""
    body = client.post("/exports", json={"task_id": seeded["task_id"], "fmt": "pptx"}).json()
    deck = Presentation(str(config.BASE_DIR / body["path"]))
    texts = [shape.text_frame.text for shape in deck.slides[0].shapes if shape.has_text_frame]
    assert any(INSIGHT in text for text in texts)
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    cells = [cell.text for row in table.rows for cell in row.cells]
    assert "地区" in cells and len(table.rows) > 1


def test_export_rejects_unknown_format_and_missing_task(client, seeded):
    """反例：格式不认识给 400，任务不存在给 404。"""
    bad = client.post("/exports", json={"task_id": seeded["task_id"], "fmt": "pdf"})
    assert bad.status_code == 400 and "只支持" in bad.json()["detail"]
    missing = client.post("/exports", json={"task_id": "t_nope", "fmt": "md"})
    assert missing.status_code == 404 and "任务不存在" in missing.json()["detail"]


def test_export_failure_leaves_no_partial_file(env, seeded, monkeypatch):
    """反例：写到一半失败不留半个文件（.part 也清掉）。"""

    def boom(doc, path):
        Path(path).write_text("半截内容", encoding="utf-8")
        raise RuntimeError("磁盘满了")

    monkeypatch.setitem(report._WRITERS, "md", boom)
    with pytest.raises(report.ReportError):
        report.export(seeded["task_id"], "md")
    leftovers = sorted(item.name for item in config.EXPORT_DIR.iterdir())
    assert leftovers == []


def test_export_without_sql_is_400(env, seeded):
    """反例：那次任务当时降级了（没有 SQL），导出给可读的 400 而不是空报告。"""
    store.insert_task({"id": "t_degraded", "dataset_id": seeded["dataset_id"], "kind": "ask", "question": QUESTION})
    with pytest.raises(report.ReportError):
        report.export("t_degraded", "md")
